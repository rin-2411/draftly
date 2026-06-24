import docx
import sys

try:
    doc = docx.Document(r'c:\Users\nadzr\Downloads\Final report FYP.docx')
    text = '\n'.join([para.text for para in doc.paragraphs])
    
    # Write to file with UTF-8 encoding
    with open('doc_content.txt', 'w', encoding='utf-8') as f:
        f.write(text)
    
    print("Document extracted successfully!")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
except Exception as e:
    print(f"Error: {e}", file=sys.stderr)
    sys.exit(1)
